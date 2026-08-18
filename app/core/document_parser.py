import io
import os
import hashlib
from typing import List, Dict, Any

import pypdf
from app.core.cache_manager import CacheManager
from loguru import logger


class DocumentParser:
    """
    文档解析器接口 (基础PDF转文本)
    """

    def __init__(self, cache_manager: CacheManager = None):
        """
        初始化文档解析器

        Args:
            cache_manager (CacheManager, optional): 缓存管理器实例
        """
        self.cache_manager = cache_manager
        logger.info("Initialized DocumentParser")

    def parse_pdf(self, file_path: str) -> str:
        """
        解析PDF文件为文本

        Args:
            file_path (str): PDF文件路径

        Returns:
            str: 提取的文本内容

        Raises:
            FileNotFoundError: 文件不存在
            ValueError: 文件不是PDF格式
            Exception: 解析失败
        """
        try:
            if not os.path.exists(file_path):
                logger.error(f"File not found: {file_path}")
                raise FileNotFoundError(f"File not found: {file_path}")

            if not file_path.lower().endswith('.pdf'):
                logger.error(f"File is not a PDF: {file_path}")
                raise ValueError(f"File is not a PDF: {file_path}")

            # 读取文件内容并生成内容哈希作为缓存键，避免路径相同但内容变化时返回旧结果
            with open(file_path, 'rb') as f:
                file_bytes = f.read()
            cache_key = f"pdf_text_{hashlib.md5(file_bytes).hexdigest()}"

            if self.cache_manager:
                cached_result = self.cache_manager.get(cache_key)
                if cached_result:
                    logger.info(f"Retrieved PDF text from cache for file: {file_path}")
                    return cached_result

            # 解析PDF
            text = ""
            pdf_reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            for page_num, page in enumerate(pdf_reader.pages):
                try:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text
                    logger.debug(f"Extracted text from page {page_num + 1}")
                except Exception as e:
                    logger.warning(f"Failed to extract text from page {page_num + 1}: {e}")
                    continue

            if self.cache_manager:
                self.cache_manager.set(cache_key, text, expire=3600)
                logger.info(f"Cached PDF text for file: {file_path}")

            logger.info(f"Parsed PDF file: {file_path}, extracted {len(text)} characters")
            return text

        except FileNotFoundError:
            raise
        except ValueError:
            raise
        except Exception as e:
            logger.error(f"Failed to parse PDF file {file_path}: {e}")
            raise Exception(f"Failed to parse PDF file {file_path}: {e}")

    def parse_docx(self, file_path: str) -> str:
        """Extract paragraph and table text from a DOCX resume."""
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover - dependency guard for old installs
            raise RuntimeError("DOCX 支持未安装，请执行 pip install -r requirements.txt") from exc
        document = Document(file_path)
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        tables = [" | ".join(cell.text.strip() for cell in row.cells) for table in document.tables for row in table.rows]
        text = "\n".join(paragraphs + tables)
        if not text.strip():
            raise ValueError("DOCX 中未提取到文本（可能是扫描图片简历）")
        return text

    def parse_multiple_pdfs(self, file_paths: List[str]) -> Dict[str, str]:
        """
        解析多个PDF文件

        Args:
            file_paths (List[str]): PDF文件路径列表

        Returns:
            Dict[str, str]: 文件路径到提取文本的映射
        """
        results = {}
        for file_path in file_paths:
            try:
                text = self.parse_pdf(file_path)
                results[file_path] = text
            except Exception as e:
                logger.error(f"Failed to parse PDF file {file_path}: {e}")
                results[file_path] = f"ERROR: {str(e)}"
        return results
